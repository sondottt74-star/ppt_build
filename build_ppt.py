"""Assemble the 3DGS paper-reading deck from assets/ into docs/3DGS论文精读.pptx."""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
A = ROOT / "assets"
OUT = ROOT.parent / "3DGS论文精读.pptx"

FONT = "PingFang SC"
LATIN = "Arial"
NAVY, TEXT, SUB = "1F3A5F", "222222", "5C677D"
LIGHT, LINE = "F3F6FA", "C9D3E3"
C1, C2, C3 = "3A86FF", "2A9D8F", "E76F51"
S_TASK, S_WHY, S_EXP = "457B9D", "E09F3E", "7B2CBF"
RED, ORANGE, YELLOW = "F4B6B6", "FAD7A0", "FFF3B0"


def I(v):
    return Emu(int(v * 914400))


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


# ------------------------------------------------------------------ text helpers
def style_run(run, size, bold=False, color=TEXT, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", LATIN), ("a:ea", FONT), ("a:cs", LATIN)):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)


def fill_para(p, text, size, color=TEXT, bold=False, accent=None):
    """`**x**` marks bold accent segments."""
    parts = text.split("**")
    for k, seg in enumerate(parts):
        if not seg:
            continue
        r = p.add_run()
        r.text = seg
        hl = k % 2 == 1
        style_run(r, size, bold=bold or hl, color=(accent or NAVY) if hl else color)


def textbox(slide, x, y, w, h, lines, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, accent=None, spacing=1.15, space_after=4):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, I(0.04))
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        fill_para(p, line, size, color, bold, accent)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, accent=None, bullet_color=None, gap=6):
    """items: str or (str, level)."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, I(0.04))
    for i, it in enumerate(items):
        text, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(I(0.26 + 0.3 * lvl))))
        pPr.set("indent", str(-int(I(0.24))))
        clr = etree.SubElement(etree.SubElement(pPr, qn("a:buClr")), qn("a:srgbClr"))
        clr.set("val", bullet_color or accent or NAVY)
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", LATIN)
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if lvl == 0 else "–")
        fill_para(p, text, size if lvl == 0 else size - 2, color if lvl == 0 else SUB, False, accent)
    return tb


def box(slide, x, y, w, h, text="", fill=LIGHT, line=None, size=14, color=TEXT, bold=False,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        accent=None, line_w=1.25, radius=0.12):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, I(0.08))
    for side in ("margin_top", "margin_bottom"):
        setattr(tf, side, I(0.04))
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.08
        fill_para(p, line, size, color, bold, accent)
    return s


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=2.0, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    t = etree.SubElement(ln, qn("a:tailEnd"))
    t.set("type", "triangle")
    t.set("w", "med")
    t.set("len", "med")
    return c


def pic(slide, name, x, y, w=None, h=None, align="center", border=None):
    """Place image fitted inside (w,h) keeping aspect ratio."""
    path = A / f"{name}.png"
    iw, ih = Image.open(path).size
    ar = iw / ih
    if w and h:
        if w / h > ar:
            nw, nh = h * ar, h
        else:
            nw, nh = w, w / ar
        ox = {"center": (w - nw) / 2, "left": 0, "right": w - nw}[align]
        oy = (h - nh) / 2
        x, y, w, h = x + ox, y + oy, nw, nh
    elif w:
        h = w / ar
    else:
        w = h * ar
    p = slide.shapes.add_picture(str(path), I(x), I(y), I(w), I(h))
    if border:
        p.line.color.rgb = rgb(border)
        p.line.width = Pt(1)
    return p, (x, y, w, h)


def caption(slide, x, y, w, text, size=11, align=PP_ALIGN.CENTER):
    return textbox(slide, x, y, w, 0.35, text, size=size, color=SUB, align=align)


# ------------------------------------------------------------------ slide chrome
page_no = [0]


CRUMBS = {
    C1: ["1 为什么选它", "2 有哪些参数", "3 形状可优化", "4 投影到屏幕", "5 实际效果"],
    C2: ["1 为什么增密", "2 怎么增删", "3 何时执行"],
    C3: ["1 像素怎么算", "2 α 混合", "3 tile 并行", "4 可微反传"],
}


def crumbs(slide, color, idx):
    steps = CRUMBS[color]
    w, gap = 1.62, 0.1
    x0 = 12.8 - len(steps) * w - (len(steps) - 1) * gap
    for k, t in enumerate(steps):
        on = k == idx
        box(slide, x0 + k * (w + gap), 0.32, w, 0.36, t, fill=color if on else "FFFFFF", line=color, size=11,
            color="FFFFFF" if on else color, bold=on, radius=0.5, line_w=1)


def new_slide(title, tag=None, tag_color=NAVY, subtitle=None, step=None):
    s = prs.slides.add_slide(BLANK)
    page_no[0] += 1
    if tag:
        box(s, 0.55, 0.32, 1.9, 0.36, tag, fill=tag_color, size=12, color="FFFFFF", bold=True, radius=0.5)
    if step is not None:
        crumbs(s, tag_color, step)
    textbox(s, 0.55, 0.72, 12.2, 0.7, title, size=28, color=NAVY, bold=True)
    if subtitle:
        textbox(s, 0.57, 1.36, 12.2, 0.4, subtitle, size=15, color=SUB)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(0.55), I(1.4 if not subtitle else 1.8), I(1.0), I(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(tag_color)
    bar.line.fill.background()
    textbox(s, 0.55, 7.08, 8, 0.3, "3D Gaussian Splatting for Real-Time Radiance Field Rendering · SIGGRAPH 2023",
            size=10, color="9AA5B8")
    textbox(s, 12.0, 7.08, 0.8, 0.3, str(page_no[0]), size=10, color="9AA5B8", align=PP_ALIGN.RIGHT)
    return s


def section_slide(num, title, sub, color):
    s = prs.slides.add_slide(BLANK)
    page_no[0] += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(NAVY)
    bg.line.fill.background()
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(0), I(3.05), I(0.35), I(1.5))
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(color)
    band.line.fill.background()
    textbox(s, 1.0, 2.2, 3, 1.3, num, size=72, color=color, bold=True)
    textbox(s, 1.0, 3.45, 11, 0.9, title, size=40, color="FFFFFF", bold=True)
    textbox(s, 1.0, 4.4, 11, 0.6, sub, size=18, color="C9D3E3")
    return s


def table(slide, x, y, w, data, col_w=None, size=11, header_rows=1, row_h=0.32,
          header_fill=NAVY, first_col_bold=True, hl_row=None):
    """Cell strings may end with ‹r›/‹o›/‹y› to mark paper's best/2nd/3rd colour."""
    rows, cols = len(data), len(data[0])
    gt = slide.shapes.add_table(rows, cols, I(x), I(y), I(w), I(row_h * rows))
    t = gt.table
    tblPr = t._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = I(cw)
    for i in range(rows):
        t.rows[i].height = I(row_h)
        for j in range(cols):
            raw = data[i][j]
            cell = t.cell(i, j)
            if raw is None:
                continue
            mark = None
            for m, col in (("‹r›", RED), ("‹o›", ORANGE), ("‹y›", YELLOW)):
                if raw.endswith(m):
                    raw, mark = raw[: -len(m)], col
            cell.margin_left = cell.margin_right = I(0.04)
            cell.margin_top = cell.margin_bottom = I(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j > 0 or i < header_rows else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = raw
            is_head = i < header_rows
            is_hl = hl_row is not None and i == hl_row
            style_run(r, size, bold=is_head or (j == 0 and first_col_bold) or is_hl,
                      color="FFFFFF" if is_head else (NAVY if is_hl and j == 0 else TEXT))
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = rgb(header_fill)
            elif mark:
                cell.fill.fore_color.rgb = rgb(mark)
            elif is_hl:
                cell.fill.fore_color.rgb = rgb("E8F0FE")
            else:
                cell.fill.fore_color.rgb = rgb("FFFFFF" if i % 2 else "F7F9FC")
    return t


def legend_marks(slide, x, y):
    for k, (col, lab) in enumerate(((RED, "第 1"), (ORANGE, "第 2"), (YELLOW, "第 3"))):
        box(slide, x + k * 1.0, y, 0.28, 0.2, fill=col, shape=MSO_SHAPE.RECTANGLE)
        textbox(slide, x + k * 1.0 + 0.32, y - 0.06, 0.7, 0.3, lab, size=10, color=SUB)


# ================================================================== slides
# 1. title --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
page_no[0] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, I(3.35))
bg.fill.solid()
bg.fill.fore_color.rgb = rgb(NAVY)
bg.line.fill.background()
textbox(s, 0.7, 0.45, 12, 0.5, "论文精读汇报", size=18, color="9EC5FE", bold=True)
textbox(s, 0.7, 0.95, 12, 1.4, "3D Gaussian Splatting for\nReal-Time Radiance Field Rendering", size=36,
        color="FFFFFF", bold=True, spacing=1.0)
textbox(s, 0.7, 2.45, 12, 0.5,
        "Bernhard Kerbl · Georgios Kopanas · Thomas Leimkühler · George Drettakis    |    SIGGRAPH 2023 (ACM TOG)",
        size=15, color="C9D3E3")
pic(s, "teaser", 0.7, 3.6, 11.9, 3.1)
caption(s, 0.7, 6.72, 11.9, "论文 Fig.1：与 Plenoxels / InstantNGP / Mip-NeRF360 对比 —— 3DGS 以 ≥100 FPS 实时渲染并达到 SOTA 画质", 12)

# 2. outline ------------------------------------------------------------
s = new_slide("汇报提纲")
items = [
    ("01", "解决的任务", "新视角合成：输入是什么，输出是什么", S_TASK),
    ("02", "为什么需要 3DGS", "NeRF 质量高但慢，点云快但质量差", S_WHY),
    ("03", "三大贡献", "① 各向异性 3D Gaussian 表示  ② 自适应密度控制  ③ 可微、可见性感知的快速光栅化器", C1),
    ("04", "实验", "数据集 · 评价指标 · 定量/定性对比 · 消融实验", S_EXP),
]
for k, (num, t, d, col) in enumerate(items):
    y = 1.85 + k * 1.22
    box(s, 0.9, y, 1.1, 0.95, num, fill=col, size=30, color="FFFFFF", bold=True)
    box(s, 2.15, y, 10.2, 0.95, "", fill=LIGHT)
    textbox(s, 2.4, y + 0.08, 9.8, 0.45, t, size=22, color=NAVY, bold=True)
    textbox(s, 2.4, y + 0.52, 9.8, 0.4, d, size=14, color=SUB)

# ================================================================== 01 task
section_slide("01", "解决的任务", "Novel View Synthesis —— 新视角合成", S_TASK)

s = new_slide("任务：新视角合成（Novel View Synthesis）", "01 任务定义", S_TASK,
              "给定同一静态场景的多张照片，生成任意新相机位姿下的照片级图像，并且要实时")
# training row
textbox(s, 0.55, 2.0, 2.0, 0.4, "训练阶段", size=16, color=S_TASK, bold=True)
for k in range(3):
    pic(s, "ds_garden", 0.6 + k * 0.16, 2.5 + k * 0.14, w=1.9, border="FFFFFF")
textbox(s, 0.55, 3.95, 2.4, 0.6, ["多视角 RGB 照片", "（静态场景）"], size=13, align=PP_ALIGN.CENTER, space_after=0)
arrow(s, 2.85, 3.15, 3.35, 3.15)
box(s, 3.4, 2.65, 1.6, 1.0, ["COLMAP", "SfM"], fill="E3ECF7", line=S_TASK, size=15, bold=True, color=NAVY)
arrow(s, 5.05, 3.15, 5.55, 3.15)
box(s, 5.6, 2.45, 2.3, 1.4, ["**相机内外参**", "焦距·主点·R·t", "**+ 稀疏点云**"], fill=LIGHT, line=LINE, size=13)
arrow(s, 7.95, 3.15, 8.45, 3.15)
box(s, 8.5, 2.65, 1.75, 1.0, ["3DGS 优化", "~30K iter"], fill=C1, size=15, bold=True, color="FFFFFF")
arrow(s, 10.3, 3.15, 10.8, 3.15)
pic(s, "aniso_real", 10.85, 2.55, w=2.0, h=1.2)
textbox(s, 10.75, 3.75, 2.2, 0.6, ["百万级 3D Gaussian", "（场景表示）"], size=13, align=PP_ALIGN.CENTER, space_after=0)
# render row
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(0.55), I(4.6), I(12.25), I(0.02))
line.fill.solid()
line.fill.fore_color.rgb = rgb(LINE)
line.line.fill.background()
textbox(s, 0.55, 4.75, 2.0, 0.4, "渲染阶段", size=16, color=S_TASK, bold=True)
box(s, 0.6, 5.3, 2.3, 1.0, ["训练好的 Gaussians", "+ 新相机位姿"], fill=LIGHT, line=LINE, size=13)
arrow(s, 2.95, 5.8, 3.65, 5.8)
box(s, 3.7, 5.3, 3.0, 1.0, ["投影 → 排序 → α 混合", "（无需反传与优化）"], fill=C1, size=14, color="FFFFFF", bold=True)
arrow(s, 6.75, 5.8, 7.45, 5.8)
pic(s, "out_garden", 7.5, 5.05, w=2.3)
textbox(s, 9.95, 5.15, 3.0, 1.3, ["**新视角 RGB 图像**", "1080p，≥ 100 FPS"], size=14)
box(s, 9.95, 6.1, 2.9, 0.5, "不输出：Mesh / 语义 / 动态", fill="FDEDEC", size=12, color="9B2C2C")

# ================================================================== 02 why
section_slide("02", "为什么需要 3DGS", "质量与速度不可兼得的矛盾", S_WHY)

s = new_slide("已有方法的矛盾：要质量就慢，要速度就差", "02 研究动机", S_WHY)
pic(s, "ray_vs_splat", 0.6, 1.65, 12.1, 3.2)
cards = [
    ("NeRF / Mip-NeRF360", ["质量 SOTA", "逐射线采样 + 查询大 MLP", "**训练 48 h，渲染 0.06 FPS**"], "FDF2E9", S_WHY),
    ("快速 NeRF（INGP / Plenoxels）", ["哈希网格 / 体素加速", "训练分钟级", "**质量下降，仍低于 20 FPS**"], LIGHT, SUB),
    ("传统点云 / 点渲染", ["显式、投影快", "点无面积 → 空洞、锯齿", "**离散，难做高质量可微优化**"], LIGHT, SUB),
]
for k, (t, lines, fill, col) in enumerate(cards):
    x = 0.6 + k * 4.12
    box(s, x, 5.0, 3.9, 1.95, "", fill=fill, line=col)
    textbox(s, x + 0.15, 5.07, 3.6, 0.4, t, size=15, color=col, bold=True)
    bullets(s, x + 0.15, 5.47, 3.6, 1.5, lines, size=13, accent="B23A48", gap=2)

s = new_slide("3DGS 的定位：显式表示 + 可微体渲染 + GPU 光栅化", "02 研究动机", S_WHY)
pic(s, "quality_speed", 0.5, 1.6, 6.9, 4.9)
caption(s, 0.5, 6.5, 6.9, "数据：论文 Table 1（Mip-NeRF360 数据集，A6000 GPU）")
box(s, 7.65, 1.7, 5.15, 1.05, ["NeRF：为每条射线**主动寻找**场景内容", "3DGS：把场景元素**主动投影**到屏幕"],
    fill="FFF6E5", line=S_WHY, size=15, align=PP_ALIGN.LEFT, accent="B5651D")
bullets(s, 7.65, 3.0, 5.2, 3.6, [
    "**像点云**：显式、局部，可直接投影、裁剪、增删",
    "**像体渲染**：连续可微，α 混合成像模型与 NeRF 相同",
    "**不在空白处采样**：只处理真正存在的 Gaussian",
    "**GPU 光栅化**：tile 并行 + 一次全局排序",
    "结果：质量 ≈ Mip-NeRF360，训练 **~40 min**，渲染 **134 FPS**",
], size=15, accent=NAVY, gap=9)

# ================================================================== 03 contributions
section_slide("03", "论文三大贡献", "① 3D Gaussian 表示   ② 自适应密度控制   ③ 可微快速光栅化器", C1)

s = new_slide("方法总览：三个贡献如何串成一个训练循环", "03 方法总览", NAVY)
pic(s, "overview", 0.6, 1.65, 12.1, 3.0)
caption(s, 0.6, 4.6, 12.1, "论文 Fig.2：从 SfM 稀疏点出发 → 3D Gaussians → 投影 → 可微 tile 光栅化 → 与 GT 比较 → 反传 + 自适应密度控制")
for k, (t, d, col) in enumerate([
    ("① 3D Gaussian 表示", "位置 · 协方差 · 不透明度 · SH 颜色", C1),
    ("② 自适应密度控制", "Clone / Split / Prune，从稀疏到稠密", C2),
    ("③ 可微快速光栅化器", "tile + 排序 + α 混合，前向/反向都快", C3),
]):
    x = 0.6 + k * 4.12
    box(s, x, 5.1, 3.9, 1.0, [f"**{t}**", d], fill="FFFFFF", line=col, size=14, accent=col, line_w=2)
textbox(s, 0.6, 6.3, 12, 0.5, "", size=12)
pic(s, "f_loss", 3.9, 6.25, w=6.0, h=0.42, align="left")
textbox(s, 0.6, 6.27, 3.3, 0.4, "训练损失（随机取一张训练图）：", size=13, color=SUB)

# ---- contribution 1
s = new_slide("贡献①：为什么用 3D Gaussian 椭球表示场景？", "03 贡献① 表示", C1, step=0)
pic(s, "primitives", 0.45, 1.75, w=6.7, h=4.4)
caption(s, 0.45, 6.25, 6.7, "示意（剖面）：用不同基元覆盖同一段表面")
box(s, 7.35, 1.7, 5.5, 1.35, ["**论文的需求（Sec.4）**",
                              "基元既要像体表示一样**可微**，又要**显式、非结构化**以便快速渲染；而输入只有**无法向的稀疏 SfM 点**"],
    fill="E8F0FE", line=C1, size=13, align=PP_ALIGN.LEFT, accent=C1)
textbox(s, 7.35, 3.2, 5.5, 0.4, "3D Gaussian 恰好同时满足：", size=16, color=NAVY, bold=True)
reasons = [
    "**连续可微**：影响随距离平滑衰减 → 参数微调时图像平滑变化，梯度能传回",
    "**不需要法向**：协方差 Σ 自带朝向，扁椭球的短轴就是隐式法向",
    "**投影有闭式解**：投影到屏幕后仍是 2D Gaussian → 可直接光栅化",
    "**显式、可增删**：能裁剪、排序、复制、删除；α 混合与 NeRF 成像模型一致",
]
for k, t in enumerate(reasons):
    y = 3.7 + k * 0.8
    box(s, 7.35, y + 0.05, 0.42, 0.42, str(k + 1), fill=C1, size=14, color="FFFFFF", bold=True, shape=MSO_SHAPE.OVAL)
    textbox(s, 7.9, y, 4.95, 0.75, t, size=13, accent=C1, spacing=1.05)

s = new_slide("一个 3D Gaussian 由哪些参数决定？", "03 贡献① 表示", C1,
              "选定椭球后需要描述：在哪（μ）· 多大多扁、朝哪（s, q → Σ）· 多透明（α）· 什么颜色（SH）", step=1)
pic(s, "gaussian_params", 0.55, 1.95, 4.6, 3.2)
pic(s, "f_gauss", 0.65, 5.35, w=4.4)
caption(s, 0.6, 5.95, 4.5, "μ 为中心，Σ 为 3×3 协方差：决定椭球的大小与朝向")
table(s, 5.5, 2.0, 7.3, [
    ["参数", "维度", "含义", "激活 / 约束"],
    ["位置 μ", "3", "椭球中心", "—"],
    ["尺度 s", "3", "三个主轴的长度", "exp(·) 保证 > 0"],
    ["旋转 q", "4", "朝向（四元数）", "归一化 → 旋转矩阵 R"],
    ["不透明度 α", "1", "透明程度", "sigmoid(·) ∈ (0,1)"],
    ["颜色 SH", "48", "随视角变化的颜色（16 系数 × RGB）", "0~3 阶球谐"],
    ["合计", "59", "每个 Gaussian 的可学习标量", ""],
], col_w=[1.4, 0.8, 3.1, 2.0], size=13, row_h=0.39, hl_row=6)
pic(s, "f_sh", 5.6, 4.85, w=4.0, h=0.8, align="left")
bullets(s, 5.5, 5.75, 7.3, 1.2, [
    "ℓ=0：与方向无关的底色；ℓ=1~3：高光等视角相关效果（每 1000 iter 开放一阶）",
    "初始化：每个 SfM 点 → 各向同性 Gaussian，半径 = 最近 3 邻居平均距离",
], size=13, gap=3)

s = new_slide("形状怎么参数化，梯度下降才不会把它“优化坏”？", "03 贡献① 表示", C1,
              "椭球的形状由 3×3 协方差 Σ 描述，但 Σ 必须半正定，直接优化 9 个数会越界 → 拆成“缩放 + 旋转”", step=2)
pic(s, "cov_decomp", 0.6, 1.95, 7.4, 2.35)
pic(s, "f_cov", 8.5, 2.5, w=4.0, h=0.7, align="left")
bullets(s, 0.7, 4.45, 6.0, 2.5, [
    "协方差必须**半正定**才有几何意义",
    "直接对 9 个元素做梯度下降 → 更新后**可能不再合法**",
    "分解后：S Sᵀ = diag(s²) ≥ 0，R 是合法旋转 → **永远合法**",
    "大小（s）与朝向（q）分开优化，**参数有几何意义**",
], size=15, accent=C1, gap=8)
box(s, 7.1, 4.4, 5.7, 2.4, "", fill=LIGHT)
textbox(s, 7.3, 4.5, 5.4, 0.4, "实际存储与优化", size=16, color=C1, bold=True)
for k, (t, d) in enumerate([("存 s（3 维）", "exp(·) 后作尺度\n保证 > 0"), ("存 q（4 维）", "归一化后\n转成旋转 R"),
                            ("用时再组合", "Σ = RSSᵀRᵀ\n梯度显式推导")]):
    box(s, 7.3 + k * 1.8, 5.05, 1.7, 1.55, [f"**{t}**", *d.split("\n")], fill="FFFFFF", line=C1, size=12, accent=C1)

s = new_slide("椭球怎么画到屏幕上：3D Gaussian → 2D Gaussian", "03 贡献① 表示", C1,
              "这正是选 Gaussian 的关键好处：投影（局部线性化）后仍是 Gaussian，屏幕椭圆由 2×2 的 Σ′ 闭式给出", step=3)
pic(s, "projection", 0.5, 1.95, 7.4, 3.75)
pic(s, "f_proj", 8.3, 2.1, w=4.4, h=0.75, align="left")
bullets(s, 8.2, 3.1, 4.7, 2.8, [
    "**W**：世界 → 相机的视角变换（旋转部分）",
    "**J**：透视投影在 Gaussian 中心处的雅可比（局部线性近似）",
    "输入 Σ（3×3）→ 输出 **Σ′（2×2）**",
    "Σ′ 决定屏幕椭圆的**长短轴、朝向与覆盖像素范围**",
], size=15, accent=C1, gap=8)
box(s, 0.6, 5.85, 12.2, 0.85, "中心 μ 用透视投影 u = fₓX/Z + cₓ 得到 μ′；形状 Σ 用局部线性化后的 J 投影 —— 即 EWA Splatting（Zwicker et al. 2001）",
    fill=LIGHT, size=14, color=SUB)

s = new_slide("实际效果：优化后的各向异性椭球自动贴合几何", "03 贡献① 表示", C1,
              "回到开头的动机：扁平 / 细长的椭球能用更少的数量贴住表面和细结构", step=4)
pic(s, "aniso_real", 0.6, 1.95, w=5.9, h=2.2, align="left")
caption(s, 0.6, 4.15, 5.9, "论文 Fig.3：左为渲染结果，右为把 Gaussian 缩小 60% 后的可视化", align=PP_ALIGN.LEFT)
bullets(s, 6.9, 2.05, 5.9, 2.2, [
    "优化后出现大量**扁平 / 细长**的 Gaussian，紧贴表面",
    "少量参数即可紧凑地覆盖复杂几何",
    "消融：Isotropic 平均 PSNR **25.23** vs Full **26.05**（30K）",
], size=16, accent=C1, gap=10)
pic(s, "abl_aniso", 0.6, 4.55, w=12.2, h=2.05)
caption(s, 0.6, 6.65, 12.2, "论文 Fig.10：各向同性 vs 各向异性（Ficus 两者都限制最多 5k 个 Gaussian）")

# ---- contribution 2
s = new_slide("贡献②：为什么需要自适应密度控制？", "03 贡献② 密度控制", C2, step=0)
bullets(s, 0.6, 1.7, 5.4, 2.4, [
    "SfM 点**稀疏**：只移动已有点，会留空洞或用大椭球糊住细节",
    "关键信号：**视空间位置梯度大** = 优化器“很想移动它” = 局部表示不够",
    "跨多个视角取**平均梯度**，超过阈值才增密：",
], size=15, accent=C2, gap=8)
pic(s, "f_grad", 0.8, 4.1, w=5.0, h=0.85, align="left")
box(s, 0.6, 5.2, 2.6, 1.55, ["**欠重建**", "几何缺失、Gaussian 太少", "→ Clone"], fill="E6F4F1", line=C2, size=14, accent=C2)
box(s, 3.4, 5.2, 2.6, 1.55, ["**过重建**", "一个大 Gaussian 盖住细节", "→ Split"], fill="E6F4F1", line=C2, size=14, accent=C2)
pic(s, "density_control", 6.3, 1.7, 6.5, 4.6)
caption(s, 6.3, 6.35, 6.5, "论文 Fig.4：上行 欠重建 → Clone；下行 过重建 → Split")

s = new_slide("Clone / Split / Prune 的决策流程", "03 贡献② 密度控制", C2, step=1)
box(s, 0.6, 1.8, 3.0, 0.8, ["每 100 iter", "（第 500 ~ 15000 iter）"], fill=NAVY, size=14, color="FFFFFF", bold=True)
arrow(s, 3.65, 2.2, 4.05, 2.2)
box(s, 4.1, 1.55, 3.7, 1.3, ["平均位置梯度", "> 0.0002 ?"], fill="FFFFFF", line=C2, size=14,
    shape=MSO_SHAPE.FLOWCHART_DECISION, bold=True, color=NAVY)
arrow(s, 7.85, 2.2, 8.3, 2.2)
textbox(s, 7.85, 1.75, 0.5, 0.3, "是", size=12, color=C2, bold=True)
box(s, 8.35, 1.55, 3.7, 1.3, ["最大尺度", "≤ 1% 场景尺度 ?"], fill="FFFFFF", line=C2, size=14,
    shape=MSO_SHAPE.FLOWCHART_DECISION, bold=True, color=NAVY)
arrow(s, 5.95, 2.88, 5.95, 3.3, color=SUB)
textbox(s, 6.05, 2.9, 0.5, 0.3, "否", size=12, color=SUB, bold=True)
box(s, 4.7, 3.35, 2.5, 0.55, "保持不变，继续优化", fill=LIGHT, size=13, color=SUB)
arrow(s, 9.6, 2.7, 8.6, 3.38, color=C2)
textbox(s, 8.0, 2.85, 0.8, 0.3, "是·小", size=12, color=C2, bold=True)
arrow(s, 10.8, 2.7, 11.5, 3.38, color=C3)
textbox(s, 11.55, 2.85, 0.8, 0.3, "否·大", size=12, color=C3, bold=True)
box(s, 7.0, 3.4, 2.75, 1.45, ["**Clone 复制**", "参数原样复制一份", "父点保留 → 数量 +1"], fill="E6F4F1", line=C2, size=13, accent=C2)
box(s, 10.0, 3.4, 2.85, 1.45, ["**Split 分裂**", "按父椭球分布采样 2 个新中心", "尺度 ÷ 1.6，删除父点"], fill="FDEDE8", line=C3, size=13, accent=C3)
box(s, 0.6, 3.3, 3.0, 1.55, ["**Prune 剪枝**（同时执行）", "α < 0.005", "或 屏幕半径 > 20 px", "或 世界尺度 > 10% 场景"],
    fill="F4F4F4", line=SUB, size=13, accent=SUB)
table(s, 0.6, 5.15, 12.25, [
    ["操作", "触发：梯度", "触发：尺寸", "新点位置", "新点尺度", "父点", "解决的问题"],
    ["Clone", "大", "小", "与父点相同（论文：沿梯度方向移动）", "不变", "保留", "小而不够：覆盖不足"],
    ["Split", "大", "大", "在父椭球内按其高斯分布采样", "÷ φ = 1.6", "删除", "大而过粗：分辨率不足"],
], col_w=[1.0, 1.15, 1.15, 3.6, 1.35, 0.9, 3.1], size=13, row_h=0.45)

s = new_slide("训练调度：增密窗口 + 周期性 opacity reset", "03 贡献② 密度控制", C2, step=2)
x0, x1, ty = 0.9, 12.4, 2.75
arrow(s, x0, ty, x1 + 0.2, ty, color=NAVY, width=3)
marks = [(0, "0"), (500, "500"), (3000, "3K"), (6000, "6K"), (9000, "9K"), (12000, "12K"), (15000, "15K"), (30000, "30K")]
pos = lambda it: x0 + (x1 - x0) * it / 30000  # noqa: E731
win = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(pos(500)), I(ty - 0.42), I(pos(15000) - pos(500)), I(0.34))
win.fill.solid()
win.fill.fore_color.rgb = rgb("CDEBE6")
win.line.fill.background()
textbox(s, pos(500), ty - 0.85, pos(15000) - pos(500), 0.4, "增密窗口：每 100 iter 执行 Clone / Split / Prune",
        size=13, color=C2, bold=True, align=PP_ALIGN.CENTER)
for it, lab in marks:
    box(s, pos(it) - 0.07, ty - 0.07, 0.14, 0.14, fill=NAVY, shape=MSO_SHAPE.OVAL)
    dx = {0: -0.25, 500: 0.2}.get(it, 0)
    textbox(s, pos(it) - 0.4 + dx, ty + 0.12, 0.8, 0.3, lab, size=12, color=NAVY, align=PP_ALIGN.CENTER, bold=True)
for it in (3000, 6000, 9000, 12000):
    box(s, pos(it) - 0.11, ty + 0.5, 0.22, 0.22, fill=C3, shape=MSO_SHAPE.OVAL)
textbox(s, pos(3000) - 0.3, ty + 0.75, 6.0, 0.35, "● opacity reset：α ← min(α, 0.01)", size=13, color=C3, bold=True)
textbox(s, pos(15000) + 0.1, ty - 0.85, 5.0, 0.4, "之后只优化参数，数量不再变化", size=13, color=SUB)
cards = [
    ("为什么要 opacity reset？", [
        "靠近相机的 floaters 一旦不透明就会“霸占”像素",
        "把 α 压低后：**有用的 Gaussian 会重新学高**，无用的停在低 α → 被 Prune",
        "相当于周期性“重新考核”每个 Gaussian",
    ], C3),
    ("其他稳定训练的技巧", [
        "**低分辨率热身**：先用 1/4 分辨率，250、500 iter 时各上采样一次",
        "**SH 逐阶开放**：先只学底色，每 1000 iter 加一阶",
        "Synthetic 场景：10 万随机点 → 迅速剪到 6~10K → 最终 20~50 万",
    ], C2),
]
for k, (t, lines, col) in enumerate(cards):
    x = 0.6 + k * 6.2
    box(s, x, 4.0, 6.0, 2.85, "", fill=LIGHT)
    textbox(s, x + 0.2, 4.1, 5.6, 0.4, t, size=16, color=col, bold=True)
    bullets(s, x + 0.2, 4.6, 5.6, 2.2, lines, size=14, accent=col, gap=6)

# ---- contribution 3
s = new_slide("贡献③：一个像素的颜色 = 多个椭球投影的有序叠加", "03 贡献③ 光栅化器", C3, step=0)
pic(s, "pixel_blend_3d", 0.4, 1.6, 8.6, 4.6)
caption(s, 0.4, 6.25, 8.6, "示意：穿过像素 p 的视线依次经过 G1、G2、G3，它们投影到图像平面后在 p 处重叠")
pic(s, "pixel_blend_screen", 9.3, 1.75, 3.6, 2.8)
caption(s, 9.3, 4.55, 3.6, "图像平面正视图：三个 2D 椭圆在 p（黑框）处叠加")
bullets(s, 9.3, 5.0, 3.7, 1.9, [
    "每个椭圆对 p 的贡献随**距离中心**衰减",
    "按**深度从近到远**依次混合",
], size=13, accent=C3, gap=4)

s = new_slide("α 混合：与 NeRF 体渲染是同一个成像模型", "03 贡献③ 光栅化器", C3, step=1)
textbox(s, 0.6, 1.65, 6.0, 0.4, "① 第 i 个椭圆在像素 p 处的有效不透明度", size=15, color=C3, bold=True)
pic(s, "f_alpha", 0.75, 2.1, w=5.9, h=0.62, align="left")
textbox(s, 0.6, 2.95, 6.0, 0.4, "② 从近到远混合（T 为剩余透射率）", size=15, color=C3, bold=True)
pic(s, "f_blend", 0.75, 3.4, w=5.9, h=0.85, align="left")
textbox(s, 0.6, 4.5, 6.0, 0.4, "对照 NeRF 体渲染：形式完全相同", size=15, color=SUB, bold=True)
pic(s, "f_nerf", 0.75, 4.95, w=5.9, h=0.6, align="left")
box(s, 0.6, 5.85, 6.0, 0.95, "区别只在样本来源：NeRF 沿射线随机采样并查询 MLP；3DGS 直接用排好序的显式 Gaussian",
    fill=LIGHT, size=13, color=SUB)
textbox(s, 7.1, 1.65, 5.8, 0.4, "手算示例：像素 p 被 3 个 Gaussian 覆盖", size=15, color=NAVY, bold=True)
table(s, 7.1, 2.15, 5.7, [
    ["", "α̃ᵢ", "Tᵢ = ∏(1-α̃ⱼ)", "权重 wᵢ = Tᵢ·α̃ᵢ"],
    ["G1（近）", "0.5", "1.0", "0.50"],
    ["G2", "0.4", "0.5", "0.20"],
    ["G3（远）", "0.8", "0.5×0.6 = 0.3", "0.24"],
    ["背景", "—", "0.3×0.2 = 0.06", "0.06"],
], col_w=[1.2, 0.9, 1.9, 1.7], size=13, row_h=0.42)
pic(s, "weights_bar", 7.2, 4.4, w=5.5)
textbox(s, 7.1, 5.85, 5.8, 1.0, ["C(p) = 0.5·c₁ + 0.2·c₂ + 0.24·c₃ + 0.06·c_bg",
                                  "顺序反过来结果就不同 → **必须按深度排序**"], size=14, accent=C3)

s = new_slide("Tile-based 光栅化：一次全局排序，tile 内并行混合", "03 贡献③ 光栅化器", C3, step=2)
steps = [
    ("视锥裁剪", "99% 置信椭球\n与视锥相交才保留"),
    ("投影", "μ→μ′，Σ→Σ′\n求覆盖哪些 tile"),
    ("按 tile 复制", "覆盖 k 个 tile\n→ k 个实例"),
    ("构造 key", "tile ID | 深度\n64 bit"),
    ("GPU 基数排序", "一次全局 Radix Sort\n（CUB）"),
    ("tile 内混合", "每 tile 一个线程块\n近→远，α≈1 提前停"),
]
for k, (t, d) in enumerate(steps):
    x = 0.55 + k * 2.1
    box(s, x, 1.75, 1.85, 1.35, [f"**{t}**", *d.split("\n")], fill="FDEDE8" if k != 5 else C3, line=C3, size=12,
        accent=C3 if k != 5 else "FFFFFF", color=TEXT if k != 5 else "FFFFFF")
    if k < 5:
        arrow(s, x + 1.87, 2.42, x + 2.08, 2.42, color=C3)
pic(s, "tiles", 0.55, 3.35, 6.0, 3.55)
textbox(s, 6.9, 3.35, 6.0, 0.4, "64 位排序 key 的布局", size=15, color=NAVY, bold=True)
box(s, 6.9, 3.85, 3.0, 0.6, "高 32 位：tile ID", fill=NAVY, size=14, color="FFFFFF", bold=True, shape=MSO_SHAPE.RECTANGLE)
box(s, 9.9, 3.85, 3.0, 0.6, "低 32 位：视空间深度", fill=C3, size=14, color="FFFFFF", bold=True, shape=MSO_SHAPE.RECTANGLE)
bullets(s, 6.9, 4.7, 6.0, 2.3, [
    "排序后：同一 tile 的实例**连续存放**，且 tile 内**按深度有序**",
    "每个 tile 只需记录 [start, end)，取出自己的列表",
    "数据批量载入 **shared memory**，像素线程协同复用",
    "代价：按中心深度排序是**近似**的 → 偶发 popping",
], size=14, accent=C3, gap=6)

s = new_slide("“可微”与“可见性感知”：反向传播同样快且不截断", "03 贡献③ 光栅化器", C3, step=3)
chain = ["图像损失", "像素颜色 C", "cᵢ , α̃ᵢ", "μ′ , Σ′", "μ, s, q, α, SH"]
for k, t in enumerate(chain):
    x = 0.6 + k * 2.5
    box(s, x, 1.75, 2.1, 0.65, t, fill=C3 if k in (0, 4) else "FDEDE8", line=C3, size=14, bold=True,
        color="FFFFFF" if k in (0, 4) else NAVY)
    if k < 4:
        arrow(s, x + 2.12, 2.07, x + 2.48, 2.07, color=C3)
textbox(s, 0.6, 2.45, 12, 0.35, "梯度链路：整个渲染过程每一步都可求导 → 照片误差可以直接更新每个 Gaussian 的参数", size=13, color=SUB)
cols = [
    ("可见性感知", [
        "按深度排序后再 α 混合 → **遮挡关系正确**",
        "Pulsar 等用与顺序无关的混合，不能正确处理遮挡",
    ]),
    ("高效反向传播", [
        "不存每像素的变长列表，**复用前向排序数组**，从后往前再遍历",
        "只存最终累积 α，反向时逐个除回 → **每像素常数内存**",
    ]),
    ("不限制梯度数量", [
        "Pulsar 只让最前面 N 个 splat 拿梯度",
        "3DGS **所有参与混合的 Gaussian 都接收梯度**",
    ]),
]
for k, (t, lines) in enumerate(cols[:2]):
    y = 3.0 + k * 1.95
    box(s, 0.6, y, 6.3, 1.8, "", fill=LIGHT)
    textbox(s, 0.8, y + 0.08, 6.0, 0.4, t, size=16, color=C3, bold=True)
    bullets(s, 0.8, y + 0.55, 6.0, 1.3, lines, size=13, accent=C3, gap=4)
t, lines = cols[2]
textbox(s, 7.2, 2.95, 5.7, 0.4, t + "：消融证据（论文 Fig.9）", size=16, color=C3, bold=True)
pic(s, "limited_bw", 7.2, 3.45, w=2.8)
pic(s, "full_bw", 10.1, 3.45, w=2.8)
caption(s, 7.2, 5.3, 2.8, "只让前 10 个拿梯度")
caption(s, 10.1, 5.3, 2.8, "完整方法")
box(s, 7.2, 5.8, 5.7, 0.95, "Truck 场景 PSNR 下降 **11 dB**（Limited-BW 13.84 vs Full 24.81）", fill="FDEDE8",
    size=14, accent=C3)

# ================================================================== 04 experiments
section_slide("04", "实验", "数据集 · 指标 · 对比结果 · 消融", S_EXP)

s = new_slide("数据集：13 个真实场景 + 8 个合成物体", "04 实验设置", S_EXP)
dsets = [
    ("Mip-NeRF360", "9 个场景", "360° 环绕拍摄的无界场景\n室外 + 室内，质量 SOTA 基准", ["ds_bicycle", "ds_room"]),
    ("Tanks & Temples", "2 个场景", "Truck、Train\n大尺度室外，前景 + 复杂远景", ["ds_truck", "ds_train"]),
    ("Deep Blending", "2 个场景", "Playroom、DrJohnson\n室内，拍摄轨迹不规则", ["ds_playroom", "ds_drjohnson"]),
    ("Synthetic NeRF", "8 个物体", "Blender 合成：相机精确、视角完整、\n白背景；用 10 万随机点初始化", ["ficus"]),
]
for k, (name, n, desc, ims) in enumerate(dsets):
    x = 0.55 + k * 3.12
    box(s, x, 1.7, 2.95, 4.15, "", fill=LIGHT)
    textbox(s, x + 0.12, 1.78, 2.75, 0.4, name, size=17, color=S_EXP, bold=True)
    textbox(s, x + 0.12, 2.18, 2.75, 0.3, n, size=13, color=SUB)
    if len(ims) == 2:
        for j, im in enumerate(ims):
            pic(s, im, x + 0.12, 2.6 + j * 0.98, w=2.71, h=0.92)
    else:
        pic(s, ims[0], x + 0.12, 2.6, w=2.71, h=1.9)
    textbox(s, x + 0.12, 4.65, 2.75, 1.1, desc.split("\n"), size=12, space_after=1)
box(s, 0.55, 6.05, 12.3, 0.8,
    ["**统一协议**：每 8 张取 1 张作为测试视角（沿用 Mip-NeRF360）；所有场景同一组超参；A6000 GPU",
     "**对比方法**：Mip-NeRF360（质量 SOTA）· InstantNGP、Plenoxels（快速方法）· 合成集另加 Mip-NeRF、Point-NeRF"],
    fill="F1E9F8", size=13, align=PP_ALIGN.LEFT, accent=S_EXP)

s = new_slide("评价指标：三个画质指标 + 三个效率指标", "04 实验设置", S_EXP)
metrics = [
    ("PSNR ↑", "峰值信噪比（dB）", "f_psnr", ["基于逐像素均方误差 MSE", "衡量**像素级保真度**", "对模糊不够敏感"]),
    ("SSIM ↑", "结构相似度（0~1）", "f_ssim", ["比较局部**亮度、对比度、结构**", "比 PSNR 更符合结构感知"]),
    ("LPIPS ↓", "学习的感知相似度", "f_lpips", ["比较 VGG/AlexNet **深度特征**距离", "最接近**人眼主观感受**"]),
]
for k, (n, sub, f, lines) in enumerate(metrics):
    x = 0.55 + k * 4.15
    box(s, x, 1.7, 3.95, 3.6, "", fill=LIGHT)
    textbox(s, x + 0.15, 1.78, 3.7, 0.45, n, size=22, color=S_EXP, bold=True)
    textbox(s, x + 0.15, 2.28, 3.7, 0.35, sub, size=13, color=SUB)
    pic(s, f, x + 0.15, 2.7, w=3.65, h=0.8)
    bullets(s, x + 0.15, 3.65, 3.7, 1.6, lines, size=13, accent=S_EXP, gap=3)
for k, (n, d) in enumerate([("Train", "训练到收敛的时间"), ("FPS", "渲染帧率，≥30 视为实时"), ("Mem", "存储训练好的模型所需空间")]):
    x = 0.55 + k * 4.15
    box(s, x, 5.5, 3.95, 0.75, [f"**{n}**：{d}"], fill="FFFFFF", line=S_EXP, size=14, accent=S_EXP)
textbox(s, 0.55, 6.4, 12.3, 0.5, "不能只看一个指标：PSNR 偏爱平滑，LPIPS 更看感知细节；还要结合速度、内存与 floaters / popping 等视觉伪影",
        size=13, color=SUB)

s = new_slide("定量结果：画质追平 SOTA，速度提升三个数量级", "04 实验结果", S_EXP, "论文 Table 1（† 为原论文数据；红/橙/黄 = 第 1/2/3 名）")
hdr = ["方法", "SSIM↑", "PSNR↑", "LPIPS↓", "Train", "FPS", "Mem"] * 1
H1 = ["数据集", "Mip-NeRF360", None, None, None, None, None, "Tanks&Temples", None, None, None, None, None,
      "Deep Blending", None, None, None, None, None]
H2 = hdr + hdr[1:] + hdr[1:]
rows_t1 = [
    ["Plenoxels", "0.626", "23.08", "0.463", "25m49s", "6.79", "2.1GB", "0.719", "21.08", "0.379", "25m5s", "13.0", "2.3GB",
     "0.795", "23.06", "0.510", "27m49s", "11.2", "2.7GB"],
    ["INGP-Base", "0.671", "25.30", "0.371", "5m37s", "11.7", "13MB", "0.723", "21.72", "0.330", "5m26s", "17.1", "13MB",
     "0.797", "23.62", "0.423", "6m31s", "3.26", "13MB"],
    ["INGP-Big", "0.699", "25.59", "0.331", "7m30s", "9.43", "48MB", "0.745", "21.92‹y›", "0.305", "6m59s", "14.4", "48MB",
     "0.817", "24.96", "0.390", "8m", "2.79", "48MB"],
    ["M-NeRF360", "0.792†‹o›", "27.69†‹r›", "0.237†‹o›", "48h", "0.06", "8.6MB", "0.759‹y›", "22.22‹o›", "0.257‹o›", "48h",
     "0.14", "8.6MB", "0.901‹o›", "29.40‹o›", "0.245‹o›", "48h", "0.09", "8.6MB"],
    ["Ours-7K", "0.770‹y›", "25.60‹y›", "0.279‹y›", "6m25s", "160", "523MB", "0.767‹o›", "21.20", "0.280‹y›", "6m55s", "197",
     "270MB", "0.875‹y›", "27.78‹y›", "0.317‹y›", "4m35s", "172", "386MB"],
    ["Ours-30K", "0.815‹r›", "27.21‹o›", "0.214‹r›", "41m33s", "134", "734MB", "0.841‹r›", "23.14‹r›", "0.183‹r›", "26m54s",
     "154", "411MB", "0.903‹r›", "29.41‹r›", "0.243‹r›", "36m2s", "137", "676MB"],
]
t = table(s, 0.45, 2.05, 12.45, [H1, H2] + rows_t1, col_w=[1.17] + [0.6, 0.66, 0.66, 0.72, 0.5, 0.6] * 3,
          size=10, header_rows=2, row_h=0.36)
for c0 in (1, 7, 13):
    t.cell(0, c0).merge(t.cell(0, c0 + 5))
legend_marks(s, 0.55, 5.1)
for k, (num, d) in enumerate([("≈", "画质与 Mip-NeRF360 相当，T&T、DB 上更好"),
                               ("48h → 40min", "训练时间缩短约 70 倍"),
                               ("0.06 → 134", "渲染 FPS 提升 2000+ 倍"),
                               ("代价", "模型 400~700 MB，远大于 NeRF 的 8.6 MB")]):
    x = 0.55 + k * 3.1
    box(s, x, 5.5, 2.95, 1.35, [f"**{num}**", d], fill=LIGHT if k < 3 else "FDEDEC", size=13,
        accent=S_EXP if k < 3 else "9B2C2C")

s = new_slide("定性对比：未见过的测试视角", "04 实验结果", S_EXP, "论文 Fig.5 节选（从上到下：Bicycle、Garden、Room、Truck）")
_, (gx, gy, gw, gh) = pic(s, "comparison_grid", 0.6, 2.35, 12.1, 4.55)
for k, lab in enumerate(["Ground Truth", "Ours (3DGS)", "Mip-NeRF360", "InstantNGP", "Plenoxels"]):
    textbox(s, gx + k * gw / 5, 1.98, gw / 5, 0.35, lab, size=14, color=S_EXP if k == 1 else NAVY, bold=True,
            align=PP_ALIGN.CENTER)

s = new_slide("合成场景 & 训练时长：7K 已可用，30K 达到最佳", "04 实验结果", S_EXP)
textbox(s, 0.55, 1.6, 7, 0.4, "Synthetic NeRF 上的 PSNR（论文 Table 2，10 万随机点初始化）", size=14, color=NAVY, bold=True)
table(s, 0.55, 2.05, 6.9, [
    ["", "Mic", "Chair", "Ship", "Mat.", "Lego", "Drums", "Ficus", "Hotdog", "Avg."],
    ["Plenoxels", "33.26", "33.98", "29.62", "29.14", "34.10", "25.35", "31.83", "36.81", "31.76"],
    ["INGP-Base", "36.22‹o›", "35.00", "31.10‹r›", "29.78‹y›", "36.39‹r›", "26.02‹y›", "33.51‹y›", "37.40‹y›", "33.18‹y›"],
    ["Mip-NeRF", "36.51‹r›", "35.14‹y›", "30.41", "30.71‹r›", "35.70‹y›", "25.48", "33.29", "37.48‹o›", "33.09"],
    ["Point-NeRF", "35.95‹y›", "35.40‹o›", "30.97‹o›", "29.61", "35.04", "26.06‹o›", "36.13‹r›", "37.30", "33.30‹o›"],
    ["Ours-30K", "35.36", "35.83‹r›", "30.80‹y›", "30.00‹o›", "35.78‹o›", "26.15‹r›", "34.87‹o›", "37.72‹r›", "33.32‹r›"],
], col_w=[1.08] + [0.646] * 9, size=10, row_h=0.36)
bullets(s, 0.55, 4.45, 6.9, 2.5, [
    "平均 PSNR **33.32**，排名第一；渲染 **180~300 FPS**",
    "合成场景相机精确、视角完整 → **随机初始化**也能成功",
    "紧凑性：与点表示方法比，用约 **1/4 点数**达到同等 PSNR",
], size=14, accent=S_EXP, gap=7)
pic(s, "abl_7k_30k", 7.8, 1.6, 5.0, 4.6)
caption(s, 7.8, 6.25, 5.0, "论文 Fig.6：7K iter（5~8 min）已有高质量；30K 主要减少背景伪影")

s = new_slide("消融实验：每个设计都在起作用", "04 消融实验", S_EXP, "论文 Table 3：PSNR，Truck / Garden / Bicycle 三个场景，5K 与 30K iter")
table(s, 0.55, 2.05, 7.4, [
    ["设置", "Truck-30K", "Garden-30K", "Bicycle-30K", "Avg-5K", "Avg-30K"],
    ["Limited-BW", "13.84", "22.88", "20.87", "19.16", "19.19"],
    ["Random Init", "18.02", "22.19", "21.05", "19.17", "20.42"],
    ["No-Split", "20.59", "26.11", "25.02", "21.50", "23.90"],
    ["No-SH", "24.39‹y›", "26.59", "25.08‹y›", "23.48‹y›", "25.35‹y›"],
    ["No-Clone", "24.82‹r›", "27.47‹o›", "25.46‹o›", "23.35", "25.91‹o›"],
    ["Isotropic", "23.89", "27.00‹y›", "24.81", "23.56‹o›", "25.23"],
    ["Full", "24.81‹o›", "27.70‹r›", "25.65‹r›", "23.90‹r›", "26.05‹r›"],
], col_w=[1.5, 1.18, 1.18, 1.18, 1.18, 1.18], size=12, row_h=0.38, hl_row=7)
chart_data = CategoryChartData()
abl = [("Limited-BW", 19.19), ("Random Init", 20.42), ("No-Split", 23.90), ("Isotropic", 25.23),
       ("No-SH", 25.35), ("No-Clone", 25.91), ("Full", 26.05)]
chart_data.categories = [a for a, _ in abl]
chart_data.add_series("Avg-30K PSNR", [v for _, v in abl])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, I(8.2), I(1.9), I(4.7), I(2.95), chart_data)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "平均 PSNR（30K）"
style_run(ch.chart_title.text_frame.paragraphs[0].runs[0], 12, bold=True, color=NAVY)
va = ch.value_axis
va.minimum_scale, va.maximum_scale = 18, 27
va.has_major_gridlines = False
va.tick_labels.font.size = Pt(10)
ch.category_axis.tick_labels.font.size = Pt(10)
plot = ch.plots[0]
plot.gap_width = 45
plot.has_data_labels = True
plot.data_labels.font.size = Pt(10)
plot.data_labels.number_format = "0.00"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for idx, pt in enumerate(plot.series[0].points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(S_EXP if idx == len(abl) - 1 else "B9A6D3")
bullets(s, 8.2, 4.95, 4.75, 2.1, [
    "**Limited-BW**：只给前 10 个 splat 梯度 → Truck 掉 11 dB",
    "**Random Init**：背景变差、floaters 增多",
    "**No-Split**：大 Gaussian 无法细化，背景变糊",
    "**No-Clone / Iso / No-SH**：各掉 0.1~0.8 dB",
], size=12, accent=S_EXP, gap=2)

s = new_slide("消融可视化：初始化与增密策略", "04 消融实验", S_EXP)
pic(s, "abl_random_sfm", 0.55, 1.65, 6.0, 4.4)
caption(s, 0.55, 6.1, 6.0, "论文 Fig.7：上 随机初始化，下 SfM 初始化 —— 主要差在背景")
pic(s, "abl_densification", 6.8, 1.65, 6.0, 4.4)
caption(s, 6.8, 6.1, 6.0, "论文 Fig.8：去掉 Split 背景模糊；去掉 Clone 细结构收敛更慢")
box(s, 0.55, 6.5, 12.25, 0.45, "结论：3DGS 的效果不是“换成 Gaussian”自然得到的 —— 表示、增密、颜色模型、渲染器共同起作用",
    fill="F1E9F8", size=13, color=S_EXP, bold=True)

s = new_slide("局限与总结", "总结", NAVY)
pic(s, "limit_mip", 0.55, 1.7, w=3.0)
pic(s, "limit_ours", 3.65, 1.7, w=3.0)
caption(s, 0.55, 3.4, 6.1, "论文 Fig.11：Mip-NeRF360 有 floaters（左）；3DGS 欠观测区域出现粗大椭球（右）")
bullets(s, 0.6, 3.95, 6.1, 3.0, [
    "**欠观测区域**：拉长 / 斑块状伪影",
    "**popping**：近似排序 + guard band 裁剪，视角变化时遮挡突变",
    "**内存大**：训练峰值可超 20 GB，模型数百 MB",
    "**依赖准确位姿**、仅限静态场景、几何不一定物理准确",
], size=14, accent="9B2C2C", gap=6)
box(s, 7.0, 1.7, 5.85, 5.2, "", fill=LIGHT)
textbox(s, 7.25, 1.82, 5.4, 0.4, "一句话总结", size=17, color=NAVY, bold=True)
flow = ["SfM 稀疏点初始化", "各向异性 3D Gaussian 表示", "投影成 2D 椭圆，按深度 α 混合",
        "照片误差可微反传，更新参数", "Clone / Split / Prune 自适应调密度", "tile 光栅化实时渲染新视角"]
cols_f = [SUB, C1, C3, C3, C2, C3]
for k, (t, col) in enumerate(zip(flow, cols_f)):
    y = 2.35 + k * 0.73
    box(s, 7.35, y, 5.15, 0.55, t, fill="FFFFFF", line=col, size=14, color=NAVY, bold=True, line_w=1.8)
    if k < len(flow) - 1:
        arrow(s, 9.92, y + 0.56, 9.92, y + 0.72, color=col)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
